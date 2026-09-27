"""Tests for app.ingestion.loaders."""

import importlib.util
import io
from pathlib import Path

import docx
import pytest
from pypdf import PdfReader, PdfWriter

from app.ingestion.document import DocumentMetadata
from app.ingestion.loaders import DocumentLoadError, extract_text

FIXTURES = Path(__file__).parent / "fixtures"

# AES support depends on an optional package the project does not install. The
# AES tests below assert whichever behaviour matches this environment, so they
# stay correct if `cryptography` is ever added.
HAS_CRYPTOGRAPHY = importlib.util.find_spec("cryptography") is not None


def make_metadata(extension: str, source_id: str = "doc") -> DocumentMetadata:
    """Minimal metadata; only ``extension`` affects loader selection."""
    return DocumentMetadata(
        source_id=source_id,
        source_type="test",
        title=f"sample{extension}",
        extension=extension,
        size_bytes=0,
    )


def build_pdf(body: str) -> bytes:
    """Construct a tiny single-page PDF containing ``body`` as real text.

    Written by hand rather than committed as a binary fixture: nothing in the
    dependency set can *generate* a text-bearing PDF (that needs reportlab),
    and a ~600 byte builder is preferable to adding a dependency or checking in
    an opaque blob.
    """
    stream = f"BT /F1 12 Tf 20 100 Td ({body}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_position = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    ).encode()
    return bytes(out)


def build_docx(paragraphs: list[str], table_rows: list[tuple[str, str]]) -> bytes:
    """Construct a .docx in memory with the given paragraphs and a 2-column table."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=2)
        for row_index, (left, right) in enumerate(table_rows):
            table.cell(row_index, 0).text = left
            table.cell(row_index, 1).text = right

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestPlainText:
    def test_utf8(self) -> None:
        data = "Hello world. Ünicode survives.".encode("utf-8")

        assert extract_text(data, make_metadata(".txt")) == "Hello world. Ünicode survives."

    def test_utf8_with_bom_strips_the_marker(self) -> None:
        data = "Hello world.".encode("utf-8-sig")

        result = extract_text(data, make_metadata(".txt"))

        assert result == "Hello world."
        assert "\ufeff" not in result, "BOM must not leak into the first chunk"

    def test_cp1252_fallback(self) -> None:
        # Valid cp1252, invalid UTF-8.
        data = b"Caf\xe9 receipt"

        assert extract_text(data, make_metadata(".txt")) == "Café receipt"

    def test_undecodable_bytes_never_raise(self) -> None:
        # 0x81 is undefined in cp1252, so only the latin-1 fallback can handle it.
        data = b"\x81\x81 trailing text"

        result = extract_text(data, make_metadata(".txt"))

        assert "trailing text" in result


class TestWhitespaceNormalisation:
    def test_runs_of_blank_lines_are_collapsed(self) -> None:
        data = b"First para.\n\n\n\n\nSecond para."

        assert extract_text(data, make_metadata(".txt")) == "First para.\n\nSecond para."

    def test_trailing_spaces_and_crlf_are_cleaned(self) -> None:
        data = b"Line one.   \r\n\r\nLine two.\t\r\n"

        assert extract_text(data, make_metadata(".txt")) == "Line one.\n\nLine two."

    def test_paragraph_breaks_are_preserved(self) -> None:
        """Chunking splits on blank lines, so they must survive extraction."""
        data = b"One.\n\nTwo.\n\nThree."

        assert extract_text(data, make_metadata(".txt")).count("\n\n") == 2


class TestDocx:
    def test_paragraphs_are_extracted(self) -> None:
        data = build_docx(["First paragraph.", "Second paragraph."], [])

        result = extract_text(data, make_metadata(".docx"))

        assert "First paragraph." in result
        assert "Second paragraph." in result

    def test_table_cells_are_extracted(self) -> None:
        """Contracts and tax documents keep the important values in tables."""
        data = build_docx(
            ["Warranty terms follow."],
            [("Coverage period", "24 months"), ("Excess", "$150")],
        )

        result = extract_text(data, make_metadata(".docx"))

        assert "24 months" in result
        assert "Coverage period" in result
        assert "$150" in result

    def test_table_row_stays_on_one_line(self) -> None:
        """A label must stay attached to its value, or retrieval loses the link."""
        data = build_docx([], [("Coverage period", "24 months")])

        result = extract_text(data, make_metadata(".docx"))

        assert "Coverage period | 24 months" in result

    def test_empty_paragraphs_are_dropped(self) -> None:
        data = build_docx(["Real text.", "", "   ", "More text."], [])

        result = extract_text(data, make_metadata(".docx"))

        assert result == "Real text.\n\nMore text."

    def test_corrupt_docx_raises(self) -> None:
        with pytest.raises(DocumentLoadError):
            extract_text(b"definitely not a docx", make_metadata(".docx"))


class TestPdf:
    def test_text_is_extracted(self) -> None:
        data = build_pdf("Hello chunk world")

        assert "Hello chunk world" in extract_text(data, make_metadata(".pdf"))

    def test_corrupt_pdf_raises(self) -> None:
        with pytest.raises(DocumentLoadError):
            extract_text(b"this is not a pdf at all", make_metadata(".pdf"))

    def test_empty_bytes_raise(self) -> None:
        with pytest.raises(DocumentLoadError):
            extract_text(b"", make_metadata(".pdf"))

    def test_encrypted_pdf_raises_rather_than_returning_nothing(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("secret")
        buffer = io.BytesIO()
        writer.write(buffer)

        with pytest.raises(DocumentLoadError, match="requires a password"):
            extract_text(buffer.getvalue(), make_metadata(".pdf"))

    def test_image_only_pdf_yields_empty_text_without_raising(self) -> None:
        """A scanned page has no extractable text; that is not an error."""
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        assert extract_text(buffer.getvalue(), make_metadata(".pdf")) == ""


def encrypt_pdf(data: bytes, *, user_password: str, algorithm: str = "RC4-128") -> bytes:
    """Re-save a PDF with encryption applied.

    RC4 needs no optional packages, so these can be generated in-test. AES
    cannot be written without ``cryptography``; those cases use the committed
    files in ``tests/fixtures``.
    """
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.encrypt(
        user_password=user_password, owner_password="owner-only", algorithm=algorithm
    )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class TestEncryptedPdf:
    """Encryption does not always mean a password is needed to read.

    Legal and government forms are often locked with an *owner* password only:
    anyone can open and read them, and the encryption merely restricts printing
    or editing. Those were rejected outright before this fix.
    """

    BODY = "Coverage period is 24 months from purchase"

    def test_owner_locked_pdf_is_read(self) -> None:
        data = encrypt_pdf(build_pdf(self.BODY), user_password="")

        assert self.BODY in extract_text(data, make_metadata(".pdf"))

    def test_pdf_needing_a_password_is_refused_with_a_clear_reason(self) -> None:
        data = encrypt_pdf(build_pdf(self.BODY), user_password="secret")

        with pytest.raises(DocumentLoadError, match="requires a password"):
            extract_text(data, make_metadata(".pdf"))

    @pytest.mark.parametrize(
        "fixture", ["aes128-owner-locked.pdf", "aes256-owner-locked.pdf"]
    )
    @pytest.mark.skipif(HAS_CRYPTOGRAPHY, reason="AES is readable when installed")
    def test_aes_without_cryptography_is_refused_not_emptied(
        self, fixture: str
    ) -> None:
        # The trap this pins: AES-128's password check needs no AES, so the
        # file "decrypts" and only fails when page text is read. Swallowing
        # that page by page returned "" with no error — the document vanished
        # from the index silently. It must be a loud, specific refusal instead.
        data = (FIXTURES / fixture).read_bytes()

        with pytest.raises(DocumentLoadError, match="cryptography"):
            extract_text(data, make_metadata(".pdf"))

    @pytest.mark.parametrize(
        "fixture", ["aes128-owner-locked.pdf", "aes256-owner-locked.pdf"]
    )
    @pytest.mark.skipif(not HAS_CRYPTOGRAPHY, reason="needs the cryptography package")
    def test_aes_is_read_when_cryptography_is_installed(self, fixture: str) -> None:
        data = (FIXTURES / fixture).read_bytes()

        assert self.BODY in extract_text(data, make_metadata(".pdf"))

    def test_fixtures_really_are_encrypted(self) -> None:
        # Guards against a fixture being regenerated without encryption, which
        # would make the AES tests above pass for the wrong reason.
        for name in ["aes128-owner-locked.pdf", "aes256-owner-locked.pdf"]:
            data = (FIXTURES / name).read_bytes()
            assert b"/Encrypt" in data, name


class TestDispatch:
    def test_unknown_extension_raises(self) -> None:
        with pytest.raises(DocumentLoadError, match="Unsupported extension"):
            extract_text(b"content", make_metadata(".xlsx"))

    def test_extension_matching_is_case_insensitive(self) -> None:
        data = "Shouty filename.".encode("utf-8")

        assert extract_text(data, make_metadata(".TXT")) == "Shouty filename."

    def test_error_message_identifies_the_document(self) -> None:
        metadata = make_metadata(".xlsx", source_id="D:/docs/budget.xlsx")

        with pytest.raises(DocumentLoadError, match="budget.xlsx"):
            extract_text(b"content", metadata)

    def test_load_error_chains_the_original_exception(self) -> None:
        """`raise ... from exc` must preserve the underlying cause."""
        with pytest.raises(DocumentLoadError) as exc_info:
            extract_text(b"not a pdf", make_metadata(".pdf"))

        assert exc_info.value.__cause__ is not None
